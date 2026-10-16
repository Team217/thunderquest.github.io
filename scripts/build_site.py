#!/usr/bin/env python3
"""Read the scoring workbook in data/, write _site/ (index.html + data.json).

Only the "Team Database" tab is read. Within it, columns are found by header
text ("Team #", "Team Name", "Round 1"...), not by position.
The Raw Performance and GP columns are ignored: each team is ranked by its
highest round score. Columns are matched by header text, so inserting/moving columns is fine.
"""
import datetime, json, os, re, shutil, sys, warnings
from pathlib import Path
from openpyxl import load_workbook

warnings.filterwarnings("ignore", message="Cannot parse header or footer")

ROOT = Path(__file__).resolve().parent.parent
DATA, SITE, OUT = ROOT / "data", ROOT / "site", ROOT / "_site"
SHEET_NAME = "Team Database"  # the only tab that is read (override with env SHEET_NAME)


def norm(v):
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else round(float(v), 2)
    try:
        return num(float(str(v).strip()))
    except ValueError:
        return None


def find_layout(ws):
    for r in range(1, min(ws.max_row, 40) + 1):
        cells = {c: norm(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 40) + 1)}
        team = next((c for c, t in cells.items() if t in ("team #", "team number", "team no.")), None)
        name = next((c for c, t in cells.items() if t == "team name"), None)
        if team and name:
            rounds = {}
            for c, t in cells.items():
                if m := re.fullmatch(r"round (\d+)", t):
                    rounds[int(m[1])] = c
            if rounds:
                return r, team, name, rounds
    return None


def read_teams(ws, layout):
    hdr, tcol, ncol, rounds = layout
    n = max(rounds)
    teams = []
    for r in range(hdr + 1, ws.max_row + 1):
        t = num(ws.cell(r, tcol).value)
        name = ws.cell(r, ncol).value
        if t is None or not name or not str(name).strip():
            continue
        rs = [num(ws.cell(r, rounds[i]).value) if i in rounds else None for i in range(1, n + 1)]
        best = max([x for x in rs if x is not None], default=None)  # highest round score; None = nothing entered
        teams.append({"num": int(t), "name": str(name).strip(), "rounds": rs, "best": best})
    return n, teams


def main():
    files = [p for p in sorted(DATA.glob("*.xls[xm]")) if not p.name.startswith("~$")]
    if not files:
        sys.exit("No .xlsx/.xlsm file found in data/")
    path = files[0]
    wb = load_workbook(path, data_only=True)  # cached values, so formulas show results
    want = os.environ.get("SHEET_NAME", SHEET_NAME)
    ws = next((w for w in wb.worksheets if norm(w.title) == norm(want)), None)
    if ws is None:
        sys.exit(f"No tab named '{want}'. Tabs in this workbook: " + ", ".join(w.title for w in wb.worksheets))
    layout = find_layout(ws)
    if not layout:
        sys.exit(f"Tab '{ws.title}' has no header row with 'Team #', 'Team Name' and 'Round 1'...")
    n, teams = read_teams(ws, layout)
    title = ws.title
    if not teams:
        sys.exit("Found headers but no teams. If the cells are formulas, open the workbook "
                 "in Excel and save it so results are stored.")

    # Competition ranking (ties share a rank: 1,1,3). A score of 0 counts and gets a rank;
    # only teams with no scores entered at all stay unranked.
    for t in teams:
        t["rank"] = (1 + sum(1 for o in teams if o["best"] is not None and o["best"] > t["best"])
                     if t["best"] is not None else None)

    payload = {
        "updated": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
        "rounds": n,
        "teams": teams,
    }
    shutil.rmtree(OUT, ignore_errors=True)
    shutil.copytree(SITE, OUT)
    blob = json.dumps(payload, separators=(",", ":"))
    (OUT / "data.json").write_text(blob, encoding="utf-8")
    idx = OUT / "index.html"
    idx.write_text(idx.read_text(encoding="utf-8").replace("/*__DATA__*/null", blob), encoding="utf-8")
    print(f"{path.name} [{title}]: {len(teams)} teams, {n} rounds -> {OUT}")


if __name__ == "__main__":
    main()